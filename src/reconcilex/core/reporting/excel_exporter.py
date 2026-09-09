"""
Production-grade Excel Reconciliation Exporter using openpyxl.
Formats professional financial multi-tab workbooks with KPI summaries and audit trails.
"""

from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reconcilex.core.models import ReconciliationReport


# Palette: Refined Slate & Zinc
HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")  # Deep slate
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="0F172A")
KPI_TITLE_FONT = Font(name="Calibri", size=9, color="64748B", bold=True)
KPI_VALUE_FONT = Font(name="Calibri", size=18, bold=True, color="0F172A")
CARD_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
ROW_ZEBRA = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
SUCCESS_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
WARNING_FILL = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
DANGER_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)


class ExcelReportExporter:
    """Generates a complete multi-tab reconciliation spreadsheet."""

    @staticmethod
    def export(report: ReconciliationReport, output_path: Path) -> Path:
        wb = openpyxl.Workbook()
        output_path.parent.mkdir(parents=True, exist_ok=True)

        # Tab 1: Executive Summary
        ws_sum = wb.active
        ws_sum.title = "Executive Summary"
        ExcelReportExporter._build_summary_sheet(ws_sum, report)

        # Tab 2: Reconciled Matches
        ws_matches = wb.create_sheet(title="Reconciled Matches")
        ExcelReportExporter._build_matches_sheet(ws_matches, report)

        # Tab 3: Unmatched Invoices
        ws_inv = wb.create_sheet(title="Unmatched Invoices")
        ExcelReportExporter._build_unmatched_invoices_sheet(ws_inv, report)

        # Tab 4: Unmatched Bank Transactions
        ws_tx = wb.create_sheet(title="Unmatched Bank Feeds")
        ExcelReportExporter._build_unmatched_transactions_sheet(ws_tx, report)

        # Tab 5: Audit Trail
        ws_audit = wb.create_sheet(title="Audit Trail")
        ExcelReportExporter._build_audit_sheet(ws_audit, report)

        wb.save(output_path)
        return output_path

    @staticmethod
    def _build_summary_sheet(ws, report: ReconciliationReport):
        ws.views.sheetView[0].showGridLines = True
        ws.cell(row=2, column=2, value="ReconcileX Financial Reconciliation Report").font = TITLE_FONT
        ws.cell(row=3, column=2, value=f"Generated: {report.generated_at.strftime('%Y-%m-%d %H:%M:%S UTC')} | Session: {report.session_id}").font = Font(size=9, color="64748B")

        # KPI Blocks
        kpis = [
            ("TOTAL INVOICES", str(report.summary.total_invoices), 5, 2),
            ("TOTAL TRANSACTIONS", str(report.summary.total_transactions), 5, 4),
            ("MATCH RATE", f"{report.summary.match_rate_percentage:.1f}%", 5, 6),
            ("MATCHED PAIRS", str(report.summary.matched_count), 8, 2),
            ("UNMATCHED INVOICES", str(report.summary.unmatched_invoices_count), 8, 4),
            ("UNMATCHED BANK FEEDS", str(report.summary.unmatched_transactions_count), 8, 6),
            ("INVOICED TOTAL", f"${report.summary.total_invoiced_amount:,.2f}", 11, 2),
            ("CLEARED BANK TOTAL", f"${report.summary.total_bank_amount:,.2f}", 11, 4),
            ("NET VARIANCE", f"${report.summary.net_variance:,.2f}", 11, 6),
        ]

        for title, val, r, c in kpis:
            # Title cell
            cell_t = ws.cell(row=r, column=c, value=title)
            cell_t.font = KPI_TITLE_FONT
            cell_t.fill = CARD_FILL
            cell_t.alignment = Alignment(horizontal="center")

            # Value cell
            cell_v = ws.cell(row=r + 1, column=c, value=val)
            cell_v.font = KPI_VALUE_FONT
            cell_v.fill = CARD_FILL
            cell_v.alignment = Alignment(horizontal="center")
            cell_v.border = THIN_BORDER

        ExcelReportExporter._auto_fit_columns(ws)

    @staticmethod
    def _build_matches_sheet(ws, report: ReconciliationReport):
        headers = [
            "Match ID", "Type", "Status", "Invoice ID(s)", "Bank Tx ID",
            "Invoice Sum ($)", "Bank Amt ($)", "Variance ($)", "Confidence", "Rule Applied", "Audit Reasoning"
        ]
        ws.append(headers)
        ExcelReportExporter._style_header(ws, len(headers))

        for idx, m in enumerate(report.matched_pairs):
            row_num = idx + 2
            inv_str = ", ".join(m.invoice_ids)
            ws.append([
                m.match_id[:8],
                m.match_type.value,
                m.match_status.value,
                inv_str,
                m.tx_id or "N/A",
                m.invoice_total,
                m.bank_amount,
                m.variance_amount,
                f"{m.confidence_score * 100:.0f}%",
                m.rule_applied,
                m.audit_reasoning
            ])

            # Apply row formatting
            fill = ROW_ZEBRA if idx % 2 == 1 else None
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=row_num, column=c)
                if fill:
                    cell.fill = fill
                cell.border = THIN_BORDER
                if c in [6, 7, 8]:
                    cell.number_format = "$#,##0.00"

        ExcelReportExporter._auto_fit_columns(ws)

    @staticmethod
    def _build_unmatched_invoices_sheet(ws, report: ReconciliationReport):
        headers = ["Doc ID", "Vendor Name", "Invoice Date", "Tax ID", "Subtotal ($)", "Tax ($)", "Total Amount ($)", "Currency", "Source Path"]
        ws.append(headers)
        ExcelReportExporter._style_header(ws, len(headers))

        for idx, inv in enumerate(report.unmatched_invoices):
            row_num = idx + 2
            ws.append([
                inv.doc_id,
                inv.vendor_name,
                inv.invoice_date.strftime("%Y-%m-%d"),
                inv.vendor_tax_id or "N/A",
                inv.subtotal or 0.0,
                inv.tax_amount or 0.0,
                inv.total_amount,
                inv.currency,
                inv.source_path
            ])
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=row_num, column=c)
                cell.border = THIN_BORDER
                cell.fill = DANGER_FILL if idx % 2 == 0 else ROW_ZEBRA
                if c in [5, 6, 7]:
                    cell.number_format = "$#,##0.00"

        ExcelReportExporter._auto_fit_columns(ws)

    @staticmethod
    def _build_unmatched_transactions_sheet(ws, report: ReconciliationReport):
        headers = ["Tx ID", "Date", "Counterparty / Narration", "Amount ($)", "Direction", "Balance ($)", "Reference"]
        ws.append(headers)
        ExcelReportExporter._style_header(ws, len(headers))

        for idx, tx in enumerate(report.unmatched_transactions):
            row_num = idx + 2
            ws.append([
                tx.tx_id,
                tx.tx_date.strftime("%Y-%m-%d"),
                tx.counterparty,
                tx.amount,
                tx.direction.value,
                tx.balance or 0.0,
                tx.reference or "N/A"
            ])
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=row_num, column=c)
                cell.border = THIN_BORDER
                cell.fill = WARNING_FILL if idx % 2 == 0 else ROW_ZEBRA
                if c in [4, 6]:
                    cell.number_format = "$#,##0.00"

        ExcelReportExporter._auto_fit_columns(ws)

    @staticmethod
    def _build_audit_sheet(ws, report: ReconciliationReport):
        headers = ["Timestamp", "Record Type", "Reference ID", "Action / Rule", "Evidence & Reasoning"]
        ws.append(headers)
        ExcelReportExporter._style_header(ws, len(headers))

        for idx, m in enumerate(report.matched_pairs):
            row_num = idx + 2
            ws.append([
                m.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
                m.match_type.value,
                f"INV:{','.join(m.invoice_ids)} <-> TX:{m.tx_id}",
                m.rule_applied,
                m.audit_reasoning
            ])
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=row_num, column=c)
                cell.border = THIN_BORDER

        ExcelReportExporter._auto_fit_columns(ws)

    @staticmethod
    def _style_header(ws, col_count: int):
        ws.views.sheetView[0].showGridLines = True
        for col in range(1, col_count + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")

    @staticmethod
    def _auto_fit_columns(ws):
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)
